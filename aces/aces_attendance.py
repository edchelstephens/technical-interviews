from pprint import pprint

from openpyxl import load_workbook, Workbook

workbook = load_workbook("Aces_Attendance.xlsx")
sheet = workbook.active

data = []

for index, row in enumerate(sheet.iter_rows(values_only=True)):
    if index == 0:
        continue

    player_name_list = [name.lower() for name in row if name is not None]
    data.extend(player_name_list)


data.sort()

attendance_dict = dict()

for name in data:
    if name in attendance_dict:
        attendance_dict[name] += 1
    else:
        attendance_dict[name] = 1

print(" ===== Attendance Dict ====== ")
pprint(attendance_dict)


output_workbook = Workbook()
sheet = output_workbook.active

sheet.title = "Aces_Attendance_Count"
sheet.append(["Name", "Attendance"])

for name, count in attendance_dict.items():
    sheet.append([name, count])

output_workbook.save("Aces_Attendance_count.xlsx")

